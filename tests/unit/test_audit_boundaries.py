from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import httpcore
import httpx
import pytest
from openpyxl import Workbook

from breachgazette.clients.base import AdapterResult, BoundedSourceClient, SourceClientError
from breachgazette.clients.deadline import DeadlineStream, DeadlineTransport
from breachgazette.clients.ipc_nsw import (
    EXPECTED_SECTORS,
    EXPECTED_TABLE_METRICS,
    NswAggregateAdapter,
    NswPublicNotificationsAdapter,
)
from breachgazette.clients.oaic import (
    EXPECTED_SHEETS,
    OaicNdbAdapter,
    oaic_metric_id,
    plan_identity_migration,
)
from breachgazette.clients.xlsx import bounded_workbook, validate_xlsx_archive
from breachgazette.compare import compare_records
from breachgazette.contracts import DateObservation, SourceNotificationRecord, SourceSnapshot
from breachgazette.pipeline import update_source
from breachgazette.privacy.audit import audit_public_value
from breachgazette.publish.builder import _build_search_assets
from breachgazette.publish.search import NORMALIZATION_VERSION, normalize_search, population_band
from breachgazette.quality.gates import DataQualityError, build_quality_report
from breachgazette.state import PrivateStateStore
from breachgazette.state.transaction import copy_candidate, state_lock
from breachgazette.utils import atomic_write_json


@pytest.mark.parametrize(
    "count,band",
    [
        (None, "not_published"),
        (0, "zero"),
        (1, "1_499"),
        (499, "1_499"),
        (500, "500_999"),
        (999, "500_999"),
        (1000, "1000_9999"),
        (9999, "1000_9999"),
        (10000, "10000_99999"),
        (99999, "10000_99999"),
        (100000, "100000_plus"),
    ],
)
def test_population_boundaries(count, band) -> None:
    assert population_band(count) == band


def test_date_routing_contains_every_searchable_year(notification_factory) -> None:
    record = notification_factory()
    record.dates.append(
        DateObservation(
            meaning="occurrence_start",
            normalized_date=date(900, 12, 20),
            origin="source_observed",
            state="present",
        )
    )
    record.dates.append(
        DateObservation(
            meaning="public_notification_date",
            normalized_date=date(2026, 1, 10),
            origin="source_observed",
            state="present",
        )
    )
    record.dates.append(
        DateObservation(
            meaning="awareness_date",
            raw_value="2027-01-01",
            normalized_date=None,
            origin="source_observed",
            state="source_conflict",
        )
    )
    unknown = notification_factory(record_id="unknown", dates=[])
    manifest, _ = _build_search_assets(
        [record, unknown], detail_ids=set(), generated_at=datetime.now(UTC)
    )
    partition = next(item for item in manifest["partitions"] if "2026" in item["id"])
    assert set(partition["years"]) == {"0900", "2025", "2026"}
    assert manifest["facet_counts"]["years"] == {
        "0900": 1,
        "2025": 1,
        "2026": 1,
        "unknown": 1,
    }
    assert manifest["query_routing"]["normalization_version"] == NORMALIZATION_VERSION


@pytest.mark.parametrize(
    "raw,normalized",
    [
        ("Straße", "strasse"),
        ("İ I \u0131", "i\u0307 i \u0131"),
        (" CAFÉ cafe\u0301 ", "café cafe\u0301"),
        ("𐐀\u2003Health\x85Group", "𐐨 health group"),
    ],
)
def test_pinned_search_normalization(raw: str, normalized: str) -> None:
    assert normalize_search(raw) == normalized


def _snapshot(count: int = 1, checksum: str = "a" * 64) -> SourceSnapshot:
    now = datetime.now(UTC)
    return SourceSnapshot(
        source_id="washington",
        retrieved_at=now,
        completed_at=now,
        revision="fixture",
        checksum_sha256=checksum,
        completeness="complete",
        records_discovered=count,
        records_accepted=count,
        records_rejected=0,
        bounded_limit=100,
        latest_attempted_update=now,
    )


@pytest.mark.parametrize(
    "method",
    [
        "initialize",
        "write_records",
        "write_snapshot",
        "append_events",
        "write_checkpoint",
        "in_progress",
        "promotion",
        "root_promotion",
        "git_promotion",
    ],
)
def test_source_generation_rolls_back_each_failed_boundary(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    notification_factory,
    method: str,
) -> None:
    root = tmp_path / "synthetic-state"
    store = PrivateStateStore(root)
    store.initialize(dataset_class="real_source_data")
    old = SourceNotificationRecord.model_validate(
        notification_factory().model_dump(exclude={"canonical_organization_id"})
    )
    store.write_records("washington", [old])
    store.write_snapshot(_snapshot())
    (root / ".git").mkdir()
    (root / ".git" / "synthetic-marker").write_text("synthetic repository metadata")
    before = {
        path: path.read_bytes()
        for path in (store.state_path("washington"), store.snapshot_path("washington"))
    }
    incoming = SourceNotificationRecord.model_validate(
        notification_factory(name="Example Corrected Health").model_dump(
            exclude={"canonical_organization_id"}
        )
    )
    result = AdapterResult(
        source_id="washington", records=[incoming], snapshot=_snapshot(checksum="b" * 64)
    )

    class Adapter:
        def collect(self, **_):
            return result

    monkeypatch.setitem(
        __import__("breachgazette.pipeline", fromlist=["ADAPTERS"]).ADAPTERS, "washington", Adapter
    )
    monkeypatch.setattr("breachgazette.pipeline.guard_source_record_count", lambda *_, **__: None)
    with monkeypatch.context() as failure:
        if method in {"promotion", "root_promotion", "git_promotion"}:
            original_rename = Path.rename

            def rename(path, target):
                if "-candidate-" in path.name and Path(target) == root and method == "promotion":
                    raise OSError("injected promotion failure")
                if method == "root_promotion" and path == root:
                    raise OSError("injected root promotion failure")
                if (
                    method == "git_promotion"
                    and path.name == ".git"
                    and "-previous" in path.parent.name
                ):
                    raise OSError("injected metadata promotion failure")
                return original_rename(path, target)

            failure.setattr(Path, "rename", rename)
        else:
            target_method = "write_checkpoint" if method == "in_progress" else method
            original = getattr(PrivateStateStore, target_method)

            def write(self, *args, **kwargs):
                if self.root != root and (
                    target_method != "write_checkpoint"
                    or args[0].status == ("in_progress" if method == "in_progress" else "complete")
                ):
                    raise OSError("injected write failure")
                return original(self, *args, **kwargs)

            failure.setattr(PrivateStateStore, target_method, write)
        with pytest.raises(OSError, match="injected"):
            update_source("washington", data_root=root)
    assert {path: path.read_bytes() for path in before} == before
    assert store.load_events() == []
    assert store.load_checkpoint("washington").status == "failed"
    assert (root / ".git" / "synthetic-marker").read_text() == "synthetic repository metadata"
    update_source("washington", data_root=root)
    first_events = store.load_events()
    assert len(first_events) == 1
    assert first_events[0].event_type == "source_record_corrected"
    first = store.load_records("washington")[0]
    assert first.local_first_observed_time == old.local_first_observed_time
    assert first.local_last_material_change_time is not None
    update_source("washington", data_root=root)
    assert store.load_events() == first_events
    assert (
        store.load_records("washington")[0].local_last_material_change_time
        == first.local_last_material_change_time
    )


def test_single_writer_lock(tmp_path: Path) -> None:
    with (
        state_lock(tmp_path / "state"),
        pytest.raises(RuntimeError, match="another writer"),
        state_lock(tmp_path / "state"),
    ):
        pytest.fail("second writer acquired the lock")


@pytest.mark.parametrize("defect", ["duplicates", "count", "source"])
def test_collection_identity_and_count_gates_fail_closed(notification_factory, defect) -> None:
    records = [notification_factory()]
    snapshots = [_snapshot()]
    if defect == "duplicates":
        snapshots.append(snapshots[0])
    elif defect == "count":
        snapshots[0].records_accepted = 2
        snapshots[0].records_discovered = 1
    else:
        records[0].source_id = "california"
    with pytest.raises(
        DataQualityError,
        match=("snapshot_counts_coherent" if defect == "count" else "source_identities_coherent"),
    ):
        build_quality_report(
            dataset_class="real_source_data",
            records_by_source={"washington": records},
            snapshots=snapshots,
        )


def test_direct_collection_rejects_inconsistent_counts_before_state_commit(
    tmp_path,
    monkeypatch,
    notification_factory,
) -> None:
    root = tmp_path / "synthetic-state"
    record = SourceNotificationRecord.model_validate(
        notification_factory().model_dump(exclude={"canonical_organization_id"})
    )
    snapshot = _snapshot()
    snapshot.records_discovered = 0

    class Adapter:
        def collect(self, **_):
            return AdapterResult(source_id="washington", records=[record], snapshot=snapshot)

    monkeypatch.setitem(
        __import__("breachgazette.pipeline", fromlist=["ADAPTERS"]).ADAPTERS, "washington", Adapter
    )
    with pytest.raises(ValueError, match="counts are inconsistent"):
        update_source("washington", data_root=root)
    store = PrivateStateStore(root)
    assert store.load_records("washington") == []
    assert store.load_snapshot("washington") is None
    assert store.load_events() == []


def test_default_deadline_transport_preserves_https_response_contract() -> None:
    response = (
        b"HTTP/1.1 200 OK\r\nContent-Length: 11\r\n"
        b'Content-Type: application/json\r\n\r\n{"ok":true}'
    )
    with httpx.Client(
        transport=DeadlineTransport(lambda: 0.5, backend=httpcore.MockBackend([response]))
    ) as client:
        assert client.get("https://synthetic.example/test").json() == {"ok": True}


def test_client_lifetime_budget_is_not_reset_for_second_url(monkeypatch) -> None:
    clock = [0.0]
    monkeypatch.setattr("breachgazette.clients.base.time.monotonic", lambda: clock[0])
    client = BoundedSourceClient(
        allowed_origins={"https://synthetic.example"},
        allowed_path_prefixes=("/",),
        max_response_bytes=100,
        total_deadline_seconds=1,
        rate_delay_seconds=0,
        transport=httpx.MockTransport(lambda _: httpx.Response(200, json={"ok": True})),
    )
    with client:
        assert client.get_json("https://synthetic.example/first")[0] == {"ok": True}
        clock[0] = 1.01
        with pytest.raises(SourceClientError, match="deadline"):
            client.get_json("https://synthetic.example/second")


def test_duplicate_state_records_fail_before_persistence(
    tmp_path: Path, notification_factory
) -> None:
    store = PrivateStateStore(tmp_path)
    with pytest.raises(ValueError, match="duplicate"):
        store.write_records("washington", [notification_factory(), notification_factory()])
    with pytest.raises(ValueError, match="source does not match"):
        store.write_records("california", [notification_factory()])
    assert not (tmp_path / "state").exists()


def test_stream_redirect_and_sleep_share_one_deadline(monkeypatch: pytest.MonkeyPatch) -> None:
    now = [0.0]
    monkeypatch.setattr("breachgazette.clients.base.time.monotonic", lambda: now[0])
    monkeypatch.setattr(
        "breachgazette.clients.base.time.sleep", lambda value: now.__setitem__(0, now[0] + value)
    )

    class Stream(httpx.SyncByteStream):
        def __iter__(self):
            for _ in range(4):
                now[0] += 0.4
                yield b"x"

    client = BoundedSourceClient(
        allowed_origins={"https://example.gov"},
        allowed_path_prefixes=("/data",),
        max_response_bytes=100,
        total_deadline_seconds=1,
        max_retries=0,
        rate_delay_seconds=0,
        transport=httpx.MockTransport(lambda _: httpx.Response(200, stream=Stream())),
    )
    with client, pytest.raises(SourceClientError):
        client.get_bytes("https://example.gov/data", accept="text/plain")
    now[0] = 0
    calls = []

    def redirects(request):
        calls.append(str(request.url))
        now[0] += 0.6
        return httpx.Response(302, headers={"Location": "/data-next"})

    with (
        BoundedSourceClient(
            allowed_origins={"https://example.gov"},
            allowed_path_prefixes=("/data",),
            max_response_bytes=100,
            total_deadline_seconds=1,
            max_retries=0,
            transport=httpx.MockTransport(redirects),
        ) as client,
        pytest.raises(SourceClientError),
    ):
        client.get_bytes("https://example.gov/data", accept="text/plain")
    assert len(calls) == 2
    now[0] = 0
    with (
        BoundedSourceClient(
            allowed_origins={"https://example.gov"},
            allowed_path_prefixes=("/data",),
            max_response_bytes=100,
            total_deadline_seconds=0.5,
            transport=httpx.MockTransport(lambda _: httpx.Response(503)),
        ) as client,
        pytest.raises(SourceClientError, match="delay"),
    ):
        client.get_bytes("https://example.gov/data", accept="text/plain")
    assert now[0] == 0


def test_each_network_read_uses_current_remaining_budget() -> None:
    timeouts = []

    class Stream:
        def read(self, size, timeout=None):
            timeouts.append(timeout)
            return b"x"

    budgets = iter([0.4, 0.1])
    stream = DeadlineStream(Stream(), lambda: next(budgets))
    stream.read(1, timeout=10)
    stream.read(1, timeout=10)
    assert timeouts == [0.4, 0.1]


def _oaic_content(value=101, *, padding=False, duplicate=False) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in EXPECTED_SHEETS:
        workbook.create_sheet(name).append(["Heading", None])
    sheet = workbook[EXPECTED_SHEETS[1]]
    if padding:
        sheet.append([None, None])
    sheet.append(["July", value])
    if duplicate:
        sheet.append(["July", value])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _oaic_result(content: bytes, period="1 July 2025 to 31 Dec 2025"):
    package = {
        "success": True,
        "result": {
            "license_id": "cc-by-4.0",
            "metadata_modified": "2026-01-01",
            "resources": [
                {
                    "id": "synthetic",
                    "format": "XLSX",
                    "name": f"NDB Data {period}",
                    "url": "https://data.gov.au/data/dataset/fixture.xlsx",
                }
            ],
        },
    }

    def response(request):
        return (
            httpx.Response(200, headers={"Content-Type": "application/json"}, json=package)
            if "/api/" in request.url.path
            else httpx.Response(200, content=content)
        )

    return OaicNdbAdapter(transport=httpx.MockTransport(response)).collect()


def test_oaic_metric_ids_are_semantic_not_values_or_offsets() -> None:
    before = _oaic_result(_oaic_content()).records[0]
    changed = _oaic_result(_oaic_content(102, padding=True)).records[0]
    new_period = _oaic_result(_oaic_content(), "1 July 2026 to 31 Dec 2026").records[0]
    assert before.source_record_id == changed.source_record_id
    assert before.source_record_id != new_period.source_record_id
    with pytest.raises(SourceClientError, match="duplicate semantic"):
        _oaic_result(_oaic_content(duplicate=True))
    args = dict(
        sheet="Individuals affected",
        period_start=date(2025, 7, 1),
        period_end=date(2025, 12, 31),
        dimension="individuals_affected",
        category="1 to 499",
        parent="Worldwide",
        unit="count",
        population_scope="Worldwide",
    )
    first = oaic_metric_id(**args)
    assert first != oaic_metric_id(
        **{**args, "parent": "Australia", "population_scope": "Australia"}
    )


def test_numeric_sector_totals_preserve_child_scope_when_blocks_move() -> None:
    def content(reverse=False):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in EXPECTED_SHEETS:
            workbook.create_sheet(name).append(["Heading", None])
        blocks = [
            [["Synthetic sector A", 30], ["Human error", 10], ["System fault", 20]],
            [["Synthetic sector B", 40], ["Human error", 10], ["System fault", 30]],
        ]
        for block in reversed(blocks) if reverse else blocks:
            for row in block:
                workbook["Top 5 sectors by source"].append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    before = _oaic_result(content()).records
    after = _oaic_result(content(True)).records
    assert len(before) == 6
    assert {record.source_record_id for record in before} == {
        record.source_record_id for record in after
    }
    children = [record for record in before if record.category == "Human error"]
    assert {record.parent_category for record in children} == {
        "Synthetic sector A",
        "Synthetic sector B",
    }
    assert len({record.source_record_id for record in children}) == 2


def test_oaic_migration_mapping_is_deterministic_and_non_mutating() -> None:
    current = _oaic_result(_oaic_content()).records[0]
    legacy = current.model_copy(deep=True)
    legacy.source_record_id = "NDB by month:2:legacy"
    legacy.normalization_version = "1.0"
    before = legacy.model_dump()
    mapping = plan_identity_migration([legacy])
    assert mapping == {legacy.source_record_id: current.source_record_id}
    assert legacy.model_dump() == before
    with pytest.raises(ValueError, match="ambiguous"):
        plan_identity_migration([legacy, legacy.model_copy()])


def test_oaic_correction_preserves_history_and_blocks_unreviewed_legacy_state(
    tmp_path,
    monkeypatch,
) -> None:
    root = tmp_path / "synthetic-state"
    results = [_oaic_result(_oaic_content()), _oaic_result(_oaic_content(102))]

    class Adapter:
        def collect(self, **_):
            return results[0]

    monkeypatch.setitem(
        __import__("breachgazette.pipeline", fromlist=["ADAPTERS"]).ADAPTERS, "oaic_ndb", Adapter
    )
    monkeypatch.setattr("breachgazette.pipeline.guard_source_record_count", lambda *_, **__: None)
    update_source("oaic_ndb", data_root=root)
    store = PrivateStateStore(root)
    first = store.load_records("oaic_ndb")[0]
    assert first.local_last_material_change_time is not None
    update_source("oaic_ndb", data_root=root)
    assert len(store.load_events()) == 1
    assert (
        store.load_records("oaic_ndb")[0].local_last_material_change_time
        == first.local_last_material_change_time
    )
    results[0] = results[1]
    update_source("oaic_ndb", data_root=root)
    corrected = store.load_records("oaic_ndb")[0]
    assert corrected.local_first_observed_time == first.local_first_observed_time
    assert [event.event_type for event in store.load_events()].count("source_record_corrected") == 1
    corrected.normalization_version = "1.0"
    store.write_records("oaic_ndb", [corrected])
    before = store.state_path("oaic_ndb").read_bytes()
    with pytest.raises(ValueError, match="migration requires review"):
        update_source("oaic_ndb", data_root=root)
    assert store.state_path("oaic_ndb").read_bytes() == before


def test_reviewed_identity_bridge_rehearsal_preserves_legacy_history(
    tmp_path,
    monkeypatch,
) -> None:
    root = tmp_path / "synthetic-legacy-state"
    original = PrivateStateStore(root)
    original.initialize(dataset_class="real_source_data")
    legacy = _oaic_result(_oaic_content()).records[0].model_copy(deep=True)
    legacy.source_record_id = "NDB by month:2:legacy"
    legacy.normalization_version = "1.0"
    legacy.local_first_observed_time = datetime(2026, 1, 1, tzinfo=UTC)
    original.write_records("oaic_ndb", [legacy])
    original.write_snapshot(_oaic_result(_oaic_content()).snapshot)
    original.append_events(
        compare_records(
            {},
            {legacy.source_record_id: legacy},
            source_id="oaic_ndb",
            previous_snapshot=None,
            current_snapshot="a" * 64,
            observed_at=legacy.local_first_observed_time,
            completeness="complete",
        )
    )
    before = {str(path): path.read_bytes() for path in root.rglob("*") if path.is_file()}
    candidate = copy_candidate(root)
    staged = PrivateStateStore(candidate)
    mapping = plan_identity_migration(staged.load_records("oaic_ndb"))
    staged.write_records(
        "oaic_ndb",
        [
            legacy.model_copy(
                update={
                    "source_record_id": mapping[legacy.source_record_id],
                    "normalization_version": "2.0",
                }
            )
        ],
    )
    atomic_write_json(
        candidate / "reviews/oaic-identity-bridge.json",
        {
            "schema_version": "1.0",
            "old_to_new": mapping,
        },
    )
    assert staged.load_events() == original.load_events()

    class Adapter:
        def collect(self, **_):
            return _oaic_result(_oaic_content(102))

    monkeypatch.setitem(
        __import__("breachgazette.pipeline", fromlist=["ADAPTERS"]).ADAPTERS, "oaic_ndb", Adapter
    )
    monkeypatch.setattr("breachgazette.pipeline.guard_source_record_count", lambda *_, **__: None)
    update_source("oaic_ndb", data_root=candidate)
    assert (
        staged.load_records("oaic_ndb")[0].local_first_observed_time
        == legacy.local_first_observed_time
    )
    events = staged.load_events()
    assert len(events) == 2
    assert events[0] == original.load_events()[0]
    assert events[1].event_type == "source_record_corrected"
    assert events[1].record_id == mapping[legacy.source_record_id]
    assert {str(path): path.read_bytes() for path in root.rglob("*") if path.is_file()} == before


def test_xlsx_dimensions_and_expanded_metadata_reject_before_iteration() -> None:
    content = _oaic_content()
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(output, "w", ZIP_DEFLATED) as changed:
        for entry in original.infolist():
            body = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet2.xml":
                body = body.replace(b'ref="A1:B2"', b'ref="A1:XFD1048576"')
            changed.writestr(entry.filename, body)
    with (
        pytest.raises(SourceClientError, match="dimensions"),
        bounded_workbook(output.getvalue(), max_rows=5000, max_columns=64),
    ):
        pytest.fail("oversized dimensions were accepted")
    tiny = BytesIO()
    with ZipFile(tiny, "w") as archive:
        archive.writestr("tiny.xml", b"x")
    malformed = bytearray(tiny.getvalue())
    offset = malformed.index(b"PK\x01\x02")
    malformed[offset + 24 : offset + 28] = (200_000_000).to_bytes(4, "little")
    with pytest.raises(SourceClientError, match="unsafe entry"):
        validate_xlsx_archive(bytes(malformed))


def test_workbook_coordinates_cannot_bypass_declared_dimensions() -> None:
    output = BytesIO()
    with (
        ZipFile(BytesIO(_oaic_content())) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as changed,
    ):
        for entry in original.infolist():
            body = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet2.xml":
                body = body.replace(b'r="A2"', b'r="XFD2"')
            changed.writestr(entry.filename, body)
    with (
        pytest.raises(SourceClientError, match="columns"),
        bounded_workbook(output.getvalue(), max_rows=5000, max_columns=64),
    ):
        pytest.fail("understated dimensions were accepted")


def test_bounded_workbook_always_closes_after_parser_failure(monkeypatch) -> None:
    closed = []

    class FakeBook:
        worksheets = ()

        def close(self):
            closed.append(True)

    monkeypatch.setattr("breachgazette.clients.xlsx.load_workbook", lambda *_, **__: FakeBook())
    with (
        pytest.raises(ValueError, match="parser failed"),
        bounded_workbook(_oaic_content(), max_rows=5000, max_columns=64),
    ):
        raise ValueError("parser failed")
    assert closed == [True]


def test_oaic_independent_rows_and_scope_blocks_can_be_reordered() -> None:
    def content(reverse):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in EXPECTED_SHEETS:
            workbook.create_sheet(name).append(["Heading", None])
        rows = [["July", 101], ["August", 102]]
        for row in reversed(rows) if reverse else rows:
            workbook["NDB by month"].append(row)
        sheet = workbook["Individuals affected"]
        sheet.append([None, None])
        sheet.append([None, None])
        blocks = [
            [["Individuals worldwide affected", None], ["1 to 499", 10]],
            [["Australians affected", None], ["1 to 499", 20]],
        ]
        for block in reversed(blocks) if reverse else blocks:
            for row in block:
                sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    before = _oaic_result(content(False)).records
    after = _oaic_result(content(True)).records
    assert {record.source_record_id for record in before} == {
        record.source_record_id for record in after
    }


@pytest.mark.parametrize(
    "row",
    ["<tr><td>1 January 2026</td></tr>", "<tr><td>1 January 2026</td><td>2025</td><td></td></tr>"],
)
def test_nsw_malformed_rows_do_not_disappear(row: str) -> None:
    header = (
        "<thead><tr><th>Agency Public Notification Date</th>"
        "<th>Agency Date of Incident</th><th>Agency Name</th></tr></thead>"
    )
    valid = "<tr><td>1 January 2026</td><td>20 December 2025</td><td>Example Agency</td></tr>"
    html = (
        f"<main><table>{header}<tbody>{valid}{row}</tbody></table>"
        f"<table>{header}<tbody></tbody></table></main>"
    )
    with pytest.raises(SourceClientError):
        NswPublicNotificationsAdapter(
            transport=httpx.MockTransport(
                lambda _: httpx.Response(200, headers={"Content-Type": "text/html"}, text=html)
            )
        ).collect()


def test_typed_asset_exception_does_not_exempt_source_text() -> None:
    digest = "0412345678abcdef" + "0" * 48
    asset = "washington-2026-001-" + digest[:16]
    manifest = {
        "schema_version": "1.0",
        "query_routing": {"algorithm": "normalized_trigram_bloom"},
        "partitions": [{"id": "washington-2026-001", "sha256": digest, "asset": asset}],
    }
    assert not audit_public_value(manifest, record_identity="manifest")
    assert audit_public_value({"asset": asset}, record_identity="source")
    manifest["partitions"][0]["asset"] = "=unsafe"
    assert audit_public_value(manifest, record_identity="manifest")


@pytest.mark.parametrize("change", ["reorder", "label", "duplicate", "unit"])
def test_nsw_metrics_use_labels_not_row_positions(
    monkeypatch: pytest.MonkeyPatch, change: str
) -> None:
    table = [["", "", "", *EXPECTED_SECTORS]]
    for index, metric in enumerate(EXPECTED_TABLE_METRICS):
        value = f"{index + 10}%" if metric.startswith("Percent") else "100"
        table.append([metric, "", "", *([value] * 5)])
    if change == "reorder":
        table[2], table[3] = table[3], table[2]
    elif change == "label":
        table[2][0] = "Changed metric"
    elif change == "duplicate":
        table[2][0] = table[3][0]
    elif change == "unit":
        table[2][3] = "100"

    class Page:
        def extract_text(self):
            return "Sector Snapshot during the reporting period"

        def extract_tables(self):
            return [table]

    class Pdf:
        pages = [Page()] * 8

        def __enter__(self):
            return self

        def __exit__(self, *_):
            pass

    monkeypatch.setattr("breachgazette.clients.ipc_nsw.pdfplumber.open", lambda _: Pdf())

    def response(request):
        if request.url.path == "/Reporting-on-the-Scheme":
            return httpx.Response(
                200,
                headers={"Content-Type": "text/html"},
                text=(
                    '<a href="/resources/mndb-scheme-data-snapshot-fixture">'
                    "MNDB Scheme Data Snapshot July to December 2025</a>"
                ),
            )
        if request.url.path.startswith("/resources/"):
            return httpx.Response(
                200,
                headers={"Content-Type": "text/html"},
                text='<a href="/media/fixture">MNDB Scheme Data Snapshot</a>',
            )
        return httpx.Response(
            200, headers={"Content-Type": "application/pdf"}, content=b"%PDF-fixture"
        )

    adapter = NswAggregateAdapter(transport=httpx.MockTransport(response))
    if change != "reorder":
        with pytest.raises(SourceClientError, match="metric"):
            adapter.collect()
    else:
        result = adapter.collect()
        assert (
            next(
                record for record in result.records if record.category == EXPECTED_TABLE_METRICS[1]
            ).value.value
            == 11
        )
        assert (
            next(
                record for record in result.records if record.category == EXPECTED_TABLE_METRICS[2]
            ).value.value
            == 12
        )


def test_nsw_duplicate_between_current_and_expired_is_rejected() -> None:
    header = (
        "<thead><tr><th>Agency Public Notification Date</th>"
        "<th>Agency Date of Incident</th><th>Agency Name</th></tr></thead>"
    )
    row = "<tr><td>1 January 2026</td><td>20 December 2025</td><td>Example Agency</td></tr>"
    html = (
        f"<main><table>{header}<tbody>{row}</tbody></table>"
        f"<table>{header}<tbody>{row}</tbody></table></main>"
    )
    with pytest.raises(SourceClientError, match="duplicate"):
        NswPublicNotificationsAdapter(
            transport=httpx.MockTransport(
                lambda _: httpx.Response(200, headers={"Content-Type": "text/html"}, text=html)
            )
        ).collect()

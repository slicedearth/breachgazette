from __future__ import annotations

import json
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from breachgazette import release
from breachgazette.cli import app
from breachgazette.clients.base import SourceClientError
from breachgazette.clients.oaic import EXPECTED_SHEETS, OaicNdbAdapter
from breachgazette.clients.oaic_review import review_legacy_workbook
from breachgazette.clients.xlsx import bounded_workbook
from breachgazette.compare import compare_records
from breachgazette.retention import create_state_backup, state_inventory
from breachgazette.state import PrivateStateStore
from breachgazette.utils import atomic_write_json, sha256_hex


def workbook_bytes(*, conflicting=False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in EXPECTED_SHEETS:
        workbook.create_sheet(name).append(["Synthetic heading", None])
    workbook["Source of breach"]["A1"] = "Synthetic causes, July-December 2025"
    for row in [
        ["System fault", 35],
        ["Synthetic repeated cause", 27],
        ["Synthetic repeated cause", 28 if conflicting else 27],
    ]:
        workbook["Source of breach"].append(row)
    sectors = workbook["Top 5 sectors by source"]
    sectors.append(["Synthetic subtitle", None])
    sectors.append(["Synthetic column heading", "Count"])
    for row in [
        ["Synthetic sector A", 30],
        ["Human error", 10],
        ["System fault", 20],
        ["Synthetic sector B", 40],
        ["Human error", 10],
        ["System fault", 30],
    ]:
        sectors.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse(content, rows=None):
    with bounded_workbook(content, max_rows=5000, max_columns=64) as workbook:
        return OaicNdbAdapter()._parse(
            workbook,
            {"id": "synthetic", "name": "NDB Data 1 July 2025 to 31 Dec 2025"},
            {"metadata_modified": "2026-01-01"},
            content,
            datetime(2026, 1, 15, tzinfo=UTC),
            identity_rows=rows,
        )


@pytest.fixture
def evidence_state(tmp_path):
    content = workbook_bytes()
    rows = {}
    result = parse(content, rows)
    current = {record.source_record_id: record for record in result.records}
    records = []
    for (sheet, offset), identity in rows.items():
        record = current[identity]
        parent = None if sheet == "Top 5 sectors by source" else record.parent_category
        legacy_id = (
            f"{sheet}:{offset}:{sha256_hex([record.category, parent, record.value.value])[:12]}"
        )
        records.append(
            record.model_copy(
                update={
                    "source_record_id": legacy_id,
                    "parent_category": parent,
                    "normalization_version": "1.0",
                }
            )
        )
    root = tmp_path / "synthetic-fixture-state"
    store = PrivateStateStore(root)
    store.initialize(dataset_class="test_fixture")
    store.write_records("oaic_ndb", records)
    snapshot = result.snapshot.model_copy(
        update={
            "records_accepted": len(records),
            "notes": result.snapshot.notes[:3],
        }
    )
    store.write_snapshot(snapshot)
    store.append_events(
        compare_records(
            {},
            {record.source_record_id: record for record in records},
            source_id="oaic_ndb",
            current_snapshot=snapshot.checksum_sha256,
            previous_snapshot=None,
            observed_at=snapshot.retrieved_at,
            completeness="complete",
        )
    )
    evidence = tmp_path / "synthetic-evidence.xlsx"
    evidence.write_bytes(content)
    return root, evidence


def test_repeated_source_cells_are_explicit_not_rejected_or_double_counted():
    result = parse(workbook_bytes())
    assert len(result.records) == 8
    assert result.snapshot.records_discovered == 9
    assert result.snapshot.records_accepted == 8
    assert result.snapshot.records_rejected == 0
    assert "Reconciled 1 exact repeated" in result.snapshot.notes[-1]
    with pytest.raises(SourceClientError, match="duplicate semantic"):
        parse(workbook_bytes(conflicting=True))


def test_repeated_cells_cannot_bypass_total_metric_bound(monkeypatch):
    monkeypatch.setattr(OaicNdbAdapter, "max_records", 8)
    with pytest.raises(SourceClientError, match="count exceeded"):
        parse(workbook_bytes())


def test_complete_review_repairs_scopes_and_preserves_every_legacy_id(evidence_state):
    root, evidence = evidence_state
    records = PrivateStateStore(root).load_records("oaic_ndb")
    before = state_inventory(root)
    mapping, parents, representatives, checksum = review_legacy_workbook(records, evidence)
    assert len(mapping) == 9 and len(representatives) == 8 and len(parents) == 4
    assert set(mapping) == {record.source_record_id for record in records}
    assert checksum == sha256_hex(evidence.read_bytes())
    assert state_inventory(root) == before


def test_review_heading_is_bounded_before_period_matching(evidence_state):
    root, evidence = evidence_state
    workbook = load_workbook(BytesIO(evidence.read_bytes()))
    workbook["Source of breach"]["A1"] = "x" * 501
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    evidence.write_bytes(content)
    records = PrivateStateStore(root).load_records("oaic_ndb")
    for record in records:
        record.source_checksum = sha256_hex(content)
    with pytest.raises(ValueError, match="500-character"):
        review_legacy_workbook(records, evidence)


@pytest.mark.parametrize(
    "mutation", ["checksum", "missing", "value", "scope", "parent", "time", "period"]
)
def test_evidence_review_rejects_unproved_transitions(evidence_state, mutation):
    root, evidence = evidence_state
    records = PrivateStateStore(root).load_records("oaic_ndb")
    if mutation == "checksum":
        evidence.write_bytes(b"synthetic altered evidence")
    elif mutation == "missing":
        records.pop()
    elif mutation == "period":
        for record in records:
            record.reporting_period_start = record.reporting_period_start.replace(year=2026)
            record.reporting_period_end = record.reporting_period_end.replace(year=2026)
    else:
        index = next(i for i, r in enumerate(records) if r.category == "Human error")
        if mutation == "time":
            index = next(
                i for i, r in enumerate(records) if r.category == "Synthetic repeated cause"
            )
            records[index].local_first_observed_time -= timedelta(days=1)
        elif mutation == "scope":
            records[index].population_scope = "Synthetic unsupported scope"
        elif mutation == "value":
            records[index].value.value = 11
        else:
            records[index].parent_category = "Synthetic unsupported parent"
        if mutation in {"value", "parent"}:
            sheet, offset, _ = records[index].source_record_id.rsplit(":", 2)
            signature = sha256_hex(
                [
                    records[index].category,
                    records[index].parent_category,
                    records[index].value.value,
                ]
            )[:12]
            records[index].source_record_id = f"{sheet}:{offset}:{signature}"
    with pytest.raises(ValueError):
        review_legacy_workbook(records, evidence)


def make_plan(root, evidence, path):
    release.plan_oaic_migration(data_root=root, output=path, workbook=evidence)
    payload = json.loads(path.read_bytes())
    payload["approved"] = True
    atomic_write_json(path, payload)
    return payload


def test_reviewed_transition_preserves_history_and_archives_original_snapshot(
    evidence_state, tmp_path
):
    root, evidence = evidence_state
    store = PrivateStateStore(root)
    before = state_inventory(root)["checksums"]
    original_snapshot = store.load_snapshot("oaic_ndb")
    original_records = store.load_records("oaic_ndb")
    plan = tmp_path / "private-reviewed-plan.json"
    payload = make_plan(root, evidence, plan)
    assert payload["schema_version"] == "1.1"
    with pytest.raises(ValueError, match="exact workbook"):
        release.migrate_oaic_identities(data_root=root, plan_path=plan)
    preview = release.migrate_oaic_identities(data_root=root, plan_path=plan, workbook=evidence)
    assert not preview["applied"] and state_inventory(root)["checksums"] == before
    backup = tmp_path / "private-backup.zip"
    checksum = create_state_backup(root, backup)["checksum_sha256"]
    result = release.migrate_oaic_identities(
        data_root=root,
        plan_path=plan,
        workbook=evidence,
        backup=backup,
        expected_backup_sha256=checksum,
        apply=True,
    )
    assert result["applied"] and result["record_count"] == 8
    after = state_inventory(root)["checksums"]
    assert {key for key in set(before) | set(after) if before.get(key) != after.get(key)} == {
        "state/oaic_ndb.json",
        "manifests/source-manifests/oaic_ndb.json",
        "reviews/oaic-identity-bridge.json",
    }
    bridge = json.loads((root / "reviews/oaic-identity-bridge.json").read_bytes())
    assert bridge["original_snapshot"] == original_snapshot.model_dump(mode="json")
    current_snapshot = store.load_snapshot("oaic_ndb")
    assert current_snapshot.records_accepted == 8
    assert current_snapshot.completed_at == original_snapshot.completed_at
    current = store.load_records("oaic_ndb")
    old_by_id = {record.source_record_id: record for record in original_records}
    for record in current:
        old = old_by_id[payload["representatives"][record.source_record_id]]
        assert record.model_dump(
            exclude={"source_record_id", "normalization_version", "parent_category"}
        ) == old.model_dump(
            exclude={"source_record_id", "normalization_version", "parent_category"}
        )
        assert record.parent_category == payload["context_updates"].get(
            old.source_record_id, old.parent_category
        )
    parsed = parse(evidence.read_bytes())
    assert not compare_records(
        {record.source_record_id: record for record in current},
        {record.source_record_id: record for record in parsed.records},
        source_id="oaic_ndb",
        current_snapshot=parsed.snapshot.checksum_sha256,
        previous_snapshot=original_snapshot.checksum_sha256,
        observed_at=parsed.snapshot.retrieved_at,
        completeness="complete",
    )


@pytest.mark.parametrize(
    "field", ["context_updates", "representatives", "evidence_checksum_sha256"]
)
def test_modified_review_plans_fail_without_state_changes(evidence_state, tmp_path, field):
    root, evidence = evidence_state
    plan = tmp_path / "private-reviewed-plan.json"
    payload = make_plan(root, evidence, plan)
    if field == "evidence_checksum_sha256":
        payload[field] = "f" * 64
    else:
        key = next(iter(payload[field]))
        payload[field][key] = "synthetic-tampering"
    atomic_write_json(plan, payload)
    before = state_inventory(root)
    with pytest.raises(ValueError):
        release.migrate_oaic_identities(data_root=root, plan_path=plan, workbook=evidence)
    assert state_inventory(root) == before


def test_evidence_cli_is_offline_and_dry_run_by_default(evidence_state, tmp_path):
    root, evidence = evidence_state
    plan = tmp_path / "private-cli-plan.json"
    before = state_inventory(root)
    runner = CliRunner()
    planned = runner.invoke(
        app,
        [
            "plan-oaic-migration",
            "--data-root",
            str(root),
            "--output",
            str(plan),
            "--workbook",
            str(evidence),
            "--json",
        ],
    )
    assert planned.exit_code == 0
    assert json.loads(planned.stdout)["record_count"] == 9
    preview = runner.invoke(
        app,
        [
            "migrate-oaic-identities",
            "--data-root",
            str(root),
            "--plan",
            str(plan),
            "--workbook",
            str(evidence),
            "--json",
        ],
    )
    assert preview.exit_code == 0
    assert not json.loads(preview.stdout)["applied"]
    missing = runner.invoke(
        app,
        [
            "migrate-oaic-identities",
            "--data-root",
            str(root),
            "--plan",
            str(plan),
            "--json",
        ],
    )
    assert missing.exit_code == 2
    assert json.loads(missing.stdout)["error"] == "identity_transition_failed"
    assert str(evidence) not in missing.stdout
    assert state_inventory(root) == before


def test_snapshot_write_failure_never_changes_original_generation(
    evidence_state, tmp_path, monkeypatch
):
    root, evidence = evidence_state
    plan = tmp_path / "private-reviewed-plan.json"
    make_plan(root, evidence, plan)
    backup = tmp_path / "private-backup.zip"
    checksum = create_state_backup(root, backup)["checksum_sha256"]
    before = state_inventory(root)

    def fail(*args, **kwargs):
        raise OSError("synthetic snapshot write failure")

    monkeypatch.setattr(PrivateStateStore, "write_snapshot", fail)
    with pytest.raises(OSError):
        release.migrate_oaic_identities(
            data_root=root,
            plan_path=plan,
            workbook=evidence,
            backup=backup,
            expected_backup_sha256=checksum,
            apply=True,
        )
    assert state_inventory(root) == before

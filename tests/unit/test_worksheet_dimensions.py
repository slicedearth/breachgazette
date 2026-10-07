from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from breachgazette.clients.base import SourceClientError
from breachgazette.clients.oaic import EXPECTED_SHEETS, OaicNdbAdapter
from breachgazette.clients.xlsx import bounded_workbook
from breachgazette.monitoring import guard_source_record_count


def metric_workbook(*, sparse=False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in EXPECTED_SHEETS:
        workbook.create_sheet(name).append(["Synthetic heading", None])
    sheet = workbook["Personal information"]
    for index in range(101):
        row = index + 2 if not sparse or index < 100 else 250
        sheet.cell(row, 1, f"Synthetic category {index}")
        sheet.cell(row, 2, index + 1)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def understate(content, before, after):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as archive, ZipFile(output, "w", ZIP_DEFLATED) as changed:
        for entry in archive.infolist():
            body = archive.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet4.xml":
                assert before in body
                body = body.replace(before, after)
            changed.writestr(entry.filename, body)
    return output.getvalue()


def parse_metrics(content):
    with bounded_workbook(content, max_rows=5_000, max_columns=64) as workbook:
        return OaicNdbAdapter()._parse(
            workbook,
            {"id": "synthetic", "name": "NDB Data 1 July 2025 to 31 Dec 2025"},
            {"metadata_modified": "2026-01-01"},
            content,
            datetime(2026, 1, 15, tzinfo=UTC),
        )


@pytest.mark.parametrize("sparse", [False, True])
def test_valid_xml_extents_keep_every_metric(sparse):
    result = parse_metrics(metric_workbook(sparse=sparse))
    assert len(result.records) == result.snapshot.records_accepted == 101
    assert result.snapshot.completeness == "complete"


@pytest.mark.parametrize("sparse,dimension", [(False, b"A1:B101"), (True, b"A1:B249")])
def test_in_bound_understated_rows_cannot_pass_the_count_floor(sparse, dimension):
    # The source floor and retention fraction alone would accept the omitted metric.
    guard_source_record_count("oaic_ndb", previous_count=101, incoming_count=100)
    content = metric_workbook(sparse=sparse)
    before = b'ref="A1:B250"' if sparse else b'ref="A1:B102"'
    altered = understate(content, before, b'ref="' + dimension + b'"')
    with pytest.raises(SourceClientError, match="dimensions"):
        parse_metrics(altered)


def test_in_bound_understated_columns_cannot_hide_numeric_values():
    content = understate(metric_workbook(), b'ref="A1:B102"', b'ref="A1:A102"')
    with pytest.raises(SourceClientError, match="dimensions"):
        parse_metrics(content)

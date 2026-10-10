"""Bound archive and worksheet metadata before iterating untrusted workbooks."""

from __future__ import annotations

import re
from collections.abc import Iterator
from contextlib import contextmanager
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import iterparse  # type: ignore[import-untyped]
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from breachgazette.clients.base import SourceClientError


def validate_xlsx_archive(
    content: bytes,
    *,
    max_entries: int = 100,
    max_entry_bytes: int = 100_000_000,
    max_expanded_bytes: int = 120_000_000,
) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > max_entries:
                raise SourceClientError("workbook archive exceeded its file-count bound")
            names: set[str] = set()
            total = 0
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or "\\" in entry.filename
                    or entry.filename in names
                    or entry.flag_bits & 1
                    or (entry.external_attr >> 16) & 0o170000 == 0o120000
                    or entry.file_size > max_entry_bytes
                ):
                    raise SourceClientError("workbook archive contained an unsafe entry")
                names.add(entry.filename)
                total += entry.file_size
            if total > max_expanded_bytes:
                raise SourceClientError("workbook archive exceeded its expanded-size bound")
    except BadZipFile as exc:
        raise SourceClientError("workbook was not a valid XLSX archive") from exc


@contextmanager
def bounded_workbook(
    content: bytes,
    *,
    max_rows: int,
    max_columns: int,
    max_entry_bytes: int = 100_000_000,
    max_expanded_bytes: int = 120_000_000,
) -> Iterator[Any]:
    validate_xlsx_archive(
        content, max_entry_bytes=max_entry_bytes, max_expanded_bytes=max_expanded_bytes
    )
    # Dimensions can understate the real XML coordinates. Check both before parsing cells.
    with ZipFile(BytesIO(content)) as archive:
        for name in archive.namelist():
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            with archive.open(name) as stream:
                rows = 0
                actual_row = actual_column = 0
                declared: tuple[int, int, int, int] | None = None
                for _, element in iterparse(stream, events=("end",)):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "dimension":
                        reference = element.get("ref", "")
                        if declared is not None or not reference or len(reference) > 32:
                            raise SourceClientError("workbook dimensions were invalid")
                        try:
                            declared = range_boundaries(reference)
                        except ValueError as error:
                            raise SourceClientError("workbook dimensions were invalid") from error
                        if (
                            any(value is None or value < 1 for value in declared)
                            or declared[0] > declared[2]
                            or declared[1] > declared[3]
                            or declared[2] > max_columns
                            or declared[3] > max_rows
                        ):
                            raise SourceClientError("workbook dimensions exceeded reviewed bounds")
                    elif tag == "row":
                        rows += 1
                        index = element.get("r", str(rows))
                        if not index.isdigit() or len(index) > 7 or int(index) < 1:
                            raise SourceClientError("workbook row coordinate was invalid")
                        actual_row = max(actual_row, int(index))
                        if rows > max_rows or actual_row > max_rows:
                            raise SourceClientError("workbook rows exceeded their reviewed bound")
                    elif tag == "c":
                        coordinate = element.get("r", "")
                        match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", coordinate)
                        if match is None:
                            raise SourceClientError("workbook cell coordinate was invalid")
                        letters, row = match.groups()
                        column = 0
                        for letter in letters:
                            column = column * 26 + ord(letter) - ord("A") + 1
                        actual_row = max(actual_row, int(row))
                        actual_column = max(actual_column, column)
                        if actual_row > max_rows:
                            raise SourceClientError("workbook rows exceeded their reviewed bound")
                        if column > max_columns:
                            raise SourceClientError(
                                "workbook columns exceeded their reviewed bound"
                            )
                    element.clear()
                if declared is None or actual_row > declared[3] or actual_column > declared[2]:
                    raise SourceClientError("workbook dimensions truncated observed XML extents")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            if (
                sheet.max_row is None
                or sheet.max_column is None
                or sheet.max_row > max_rows
                or sheet.max_column > max_columns
            ):
                raise SourceClientError("workbook dimensions exceeded their reviewed bounds")
        yield workbook
    finally:
        workbook.close()
